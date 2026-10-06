# -*- coding: utf-8 -*-
"""
Informe mensual de Control de Calidad en Word (.docx).

Sigue la estructura del informe "Informe de Control de Calidad - Enero 2025":
memo, resumen de lotes, duplicados, blancos, estandares, check externo,
acumulado, comentarios, conclusiones y Anexo A con los graficos.

Los textos con numeros (cantidades, tasas de insercion, tasas de error,
clasificaciones) se completan automaticamente. Lo que requiere criterio del
geologo queda como texto resaltado en amarillo "[ ... ]" para completarlo en
Word. No depende de Qt: se puede generar desde un script.

@author: frecines
"""
import io
from datetime import datetime

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import qaqc_config as cfg
import qaqc_estandares as estandares

COLOR_CLASE = {"ACEPTABLE": RGBColor(0x00, 0x80, 0x00), "CUESTIONABLE": RGBColor(0xE6, 0x8A, 0x00),
               "INACEPTABLE": RGBColor(0xFF, 0x00, 0x00), "REFERENCIAL": RGBColor(0x1F, 0x4E, 0x9A)}
GRIS = "D9D9D9"
COLOR_HEX = {"yellow": "FFFF00", "orange": "FFA500", "lightgreen": "90EE90",
             "lightblue": "ADD8E6", "pink": "FFC0CB", "khaki": "F0E68C",
             "plum": "DDA0DD", "wheat": "F5DEB3", "green": "008000"}
DPI_FIGURAS = 120

# Criterio de los certificados de MRC (Smee & Associates Consulting Ltd.)
CRITERIO_CV = (
    "De acuerdo con los certificados de análisis de los MRC (Smee & Associates Consulting "
    "Ltd.), la media certificada y la desviación estándar entre laboratorios se obtienen de un "
    "round robin con varios laboratorios. Los elementos con desviación estándar relativa (RSD o "
    f"CV) cercana o menor a {cfg.UMBRAL_CV_ESTANDAR}% se clasifican como Certificados; entre "
    f"{cfg.UMBRAL_CV_ESTANDAR}% y {cfg.UMBRAL_CV_INDICATIVO}% como Provisionales, que deben "
    "usarse con precaución al evaluar la exactitud de un análisis individual; y mayores a "
    f"{cfg.UMBRAL_CV_INDICATIVO}% como Indicativos, que no pueden usarse para monitorear la "
    f"exactitud. Por ello, los elementos con CV mayor a {cfg.UMBRAL_CV_ESTANDAR}% se grafican y "
    "se presentan como valores referenciales, sin evaluación de sesgo.")


# ----------------------------------------------------------------- textos
def lista_es(items):
    """['Cu', 'Ag', 'Zn'] -> 'Cu, Ag y Zn'."""
    items = [str(i) for i in items]
    if not items:
        return ""
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " y " + items[-1]


def _elemento(analito):
    return str(analito).split()[0]


def _pct(valor, decimales=2):
    if pd.isna(valor):
        return "-"
    texto = f"{valor:.{decimales}f}".rstrip("0").rstrip(".")
    return texto or "0"


def texto_tasas(rep, col_pct):
    """Frase con el rango de tasas de error de un cuadro de reporte."""
    datos = rep[["Analito", col_pct]].dropna()
    if datos.empty:
        return "No hubo pares validos para evaluar."
    vmin, vmax = datos[col_pct].min(), datos[col_pct].max()
    if np.isclose(vmin, vmax):
        return f"La tasa de error fue de {_pct(vmin)}% para todos los elementos"
    con_min = lista_es(_elemento(a) for a in datos.loc[np.isclose(datos[col_pct], vmin), "Analito"])
    con_max = lista_es(_elemento(a) for a in datos.loc[np.isclose(datos[col_pct], vmax), "Analito"])
    return (f"Las tasas de error oscilaron entre {_pct(vmin)}% para {con_min} "
            f"y {_pct(vmax)}% para {con_max}")


def elementos_con_clase(rep, col_clase, clase, col_nombre="Analito", prefijo=None):
    filas = rep[rep[col_clase] == clase]
    if prefijo:
        return [f"{r[prefijo]} {_elemento(r[col_nombre])}" for _, r in filas.iterrows()]
    return [_elemento(a) for a in filas[col_nombre]]


def texto_aceptabilidad(inaceptables, sujeto="todos los elementos"):
    if not inaceptables:
        return f"aceptable para {sujeto}"
    return f"aceptable para {sujeto}, excepto {lista_es(inaceptables)}"


# ------------------------------------------------------------ utilidades docx
def _sombrear(celda, color_hex):
    tc_pr = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color_hex)
    tc_pr.append(sombra)


def _campo(parrafo, codigo):
    """Inserta un campo de Word (PAGE, NUMPAGES)."""
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, codigo), ("end", None)):
        if tipo:
            elem = OxmlElement("w:fldChar")
            elem.set(qn("w:fldCharType"), tipo)
        else:
            elem = OxmlElement("w:instrText")
            elem.set(qn("xml:space"), "preserve")
            elem.text = texto
        run._r.append(elem)


class Informe:
    def __init__(self):
        self.doc = Document()
        self.n_tabla = 0
        estilo = self.doc.styles["Normal"]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(11)
        estilo.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        estilo.paragraph_format.space_after = Pt(6)
        seccion = self.doc.sections[0]
        seccion.page_width, seccion.page_height = Cm(21.59), Cm(27.94)
        seccion.left_margin = seccion.right_margin = Cm(2.54)
        seccion.top_margin, seccion.bottom_margin = Cm(3.0), Cm(2.0)
        seccion.header.paragraphs[0].text = "Geología Mina"
        pie = seccion.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pie.add_run("Página ")
        _campo(pie, "PAGE")
        pie.add_run(" de ")
        _campo(pie, "NUMPAGES")

    # textos
    def titulo(self, texto, centrado=False, tamano=None):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(10)
        if centrado:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(texto)
        run.bold = True
        if tamano:
            run.font.size = Pt(tamano)
        return p

    def subtitulo(self, texto):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(texto)
        run.bold = True
        run.italic = True
        return p

    def parrafo(self, *partes, vineta=False):
        """partes: str (texto normal) o ('PENDIENTE', texto) para resaltar."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if vineta:
            p.paragraph_format.left_indent = Cm(0.63)
            p.paragraph_format.first_line_indent = Cm(-0.63)
            p.add_run("-\t")
        for parte in partes:
            if isinstance(parte, tuple):
                run = p.add_run(parte[1])
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            else:
                p.add_run(parte)
        return p

    def pendiente(self, texto):
        return self.parrafo(("PENDIENTE", f"[{texto}]"))

    def leyenda(self, texto):
        self.n_tabla += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"Tabla {self.n_tabla}: {texto}")
        run.bold = True
        return self.n_tabla

    # tablas
    def tabla(self, df, col_color=None, colores=None, gris_primera=True, tamano=None,
              negrita_filas=(), gris_filas=()):
        filas, columnas = df.shape
        tamano = tamano or (7 if columnas > 8 else 8)
        t = self.doc.add_table(rows=filas + 1, cols=columnas)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, nombre in enumerate(df.columns):
            celda = t.cell(0, j)
            celda.text = str(nombre).replace("\n", " ")
            _sombrear(celda, GRIS)
            for run in celda.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(tamano)
            celda.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        clases = {j for j, c in enumerate(df.columns) if "Clase" in str(c)}
        for i in range(filas):
            for j in range(columnas):
                valor = df.iat[i, j]
                texto = "" if not isinstance(valor, str) and pd.isna(valor) else str(valor)
                celda = t.cell(i + 1, j)
                celda.text = texto
                run = celda.paragraphs[0].runs[0] if celda.paragraphs[0].runs else None
                if run is not None:
                    run.font.size = Pt(tamano)
                    if j in clases and texto in COLOR_CLASE:
                        run.font.color.rgb = COLOR_CLASE[texto]
                    if i in negrita_filas:
                        run.bold = True
                if j == col_color and colores:
                    _sombrear(celda, COLOR_HEX.get(colores.get(texto, ""), GRIS))
                elif (j == 0 and gris_primera) or i in gris_filas:
                    _sombrear(celda, GRIS)
        self._ajustar_anchos(t, df, tamano)
        self.doc.add_paragraph()
        return t

    @staticmethod
    def _ajustar_anchos(tabla, df, tamano, ancho_total=16.5):
        """Ancho de columnas segun la palabra mas larga de cada columna, para que
        'ACEPTABLE' o 'CHNB13' no se partan en dos lineas."""
        margenes = OxmlElement("w:tblCellMar")          # relleno de celda 0.1 cm
        for lado in ("left", "right"):
            m = OxmlElement(f"w:{lado}")
            m.set(qn("w:w"), "57")
            m.set(qn("w:type"), "dxa")
            margenes.append(m)
        tabla._tbl.tblPr.append(margenes)

        def palabra_larga(valor):  # las mayusculas ocupan ~1.3 veces mas
            return max((sum(1.65 if c in "MWmw" else 1.35 if c.isupper() else 1.0 for c in p)
                        for p in str(valor).split()), default=1)

        cm_por_caracter = 0.0215 * tamano
        anchos = []
        for j, nombre in enumerate(df.columns):
            valores = [v for v in df.iloc[:, j] if not (not isinstance(v, str) and pd.isna(v))]
            contenido = max((palabra_larga(v) for v in valores), default=1)
            anchos.append(max(palabra_larga(nombre), contenido) * cm_por_caracter + 0.25)
        escala = min(1.0, ancho_total / sum(anchos))
        tabla.autofit = False
        for j, ancho in enumerate(anchos):
            tabla.columns[j].width = Cm(ancho * escala)      # grilla (w:gridCol)
            for celda in tabla.columns[j].cells:
                celda.width = Cm(ancho * escala)

    # figuras
    @staticmethod
    def _png(tamano, dibujar):
        from matplotlib.figure import Figure
        fig = Figure(figsize=tamano)
        fig.estatica = True          # sin tooltips (no hay ventana)
        dibujar(fig)
        buffer = io.BytesIO()
        fig.savefig(buffer, format="png", dpi=DPI_FIGURAS, bbox_inches="tight")
        buffer.seek(0)
        return buffer

    def figuras(self, graficos, columnas=2, ancho_cm=None, avance=None):
        """graficos: [(pestana, archivo, figsize, dibujar)] en una grilla."""
        if not graficos:
            return
        ancho_cm = ancho_cm or (16.0 / columnas)
        t = self.doc.add_table(rows=(len(graficos) + columnas - 1) // columnas, cols=columnas)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, (_, nombre, tamano, dibujar) in enumerate(graficos):
            if avance:
                avance(nombre)
            celda = t.cell(k // columnas, k % columnas)
            parrafo = celda.paragraphs[0]
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parrafo.add_run().add_picture(self._png(tamano, dibujar), width=Cm(ancho_cm))
        self.doc.add_paragraph()

    def salto_pagina(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def memo(self, para, de, fecha, asunto, anexos):
        t = self.doc.add_table(rows=6, cols=2)
        filas = [("Memo", ""), ("A:", para), ("De:", de), ("Fecha:", fecha),
                 ("Asunto:", asunto), ("Anexos:", anexos)]
        for i, (etiqueta, valor) in enumerate(filas):
            a, b = t.cell(i, 0), t.cell(i, 1)
            a.width, b.width = Cm(3), Cm(13)
            run = a.paragraphs[0].add_run(etiqueta)
            run.bold = True
            if i == 0:
                run.font.size = Pt(20)
            else:
                b.paragraphs[0].add_run(valor).bold = i == 4
        borde = OxmlElement("w:pBdr")
        inferior = OxmlElement("w:bottom")
        for clave, valor in (("val", "single"), ("sz", "12"), ("space", "1"), ("color", "000000")):
            inferior.set(qn(f"w:{clave}"), valor)
        borde.append(inferior)
        separador = self.doc.add_paragraph()
        separador._p.get_or_add_pPr().append(borde)

    def guardar(self, ruta):
        self.doc.save(ruta)


# --------------------------------------------------------------- contenido
def tabla_estandares_usados(conf_std):
    """Tabla de MRCs utilizados (nombre, elemento, MH, BV, CV, limites, observaciones)."""
    filas = []
    for conf in conf_std.values():
        for e in cfg.ELEMENTOS_STD:
            mh, de = conf["MH"][e], conf["DE"][e]
            bv, cv = conf["BV"].get(e), conf.get("CV", {}).get(e)
            categoria = estandares.categoria_cv(cv)
            if bv is not None and estandares.es_referencial(conf, e):
                obs = f"{categoria} (CV > {cfg.UMBRAL_CV_ESTANDAR}%): valor referencial."
            elif bv is None:
                obs = "Sin valor certificado (provisional)."
            else:
                obs = categoria or ""
            filas.append([conf["codigo"], cfg.ANALITO[e], round(mh, 4),
                          "" if bv is None else round(bv, 4),
                          "" if cv is None else round(cv, 2),
                          round(mh - 2 * de, 4), round(mh + 2 * de, 4), obs])
    return pd.DataFrame(filas, columns=["Nombre STD", "Elemento", "Media histórica (MH)",
                                        "Media certificada (BV)", "CV BV (%)", "MH-2STDEV",
                                        "MH+2STDEV", "Observaciones"])


def tabla_limites_deteccion():
    unidades = {"AG": "g/tn"}
    return pd.DataFrame([[cfg.ETIQUETA[e], unidades.get(e, "%"), cfg.LD_LABORATORIO[e],
                          cfg.LD_DUPLICADO[e]] for e in cfg.ELEMENTOS],
                        columns=["Elemento", "Unidades", "Lím. Det. (declarado por el laboratorio)",
                                 "Lím. Det. (determinado de duplicados de pulpa)"])


def generar(ruta, datos, anexo, acumulado=(), avance=None):
    """Genera el informe.

    datos: dict con
        para, de, fecha, mes, anio, periodo_desde, periodo_hasta,
        res (ResultadosQAQC), ratio (DataFrame de calc.ratio_insercion),
        reporte_rma (DataFrame o None), colores (codigo STD -> color).
    anexo: [(titulo_seccion, [graficos])] para el Anexo A.
    acumulado: graficos historicos de estandares para "Analisis del acumulado".
    """
    res, rep = datos["res"], datos["res"].reportes
    ratio = datos["ratio"].set_index("Año/Mes")
    col_n, col_ri = ratio.columns[0], ratio.columns[1]
    periodo = f"{datos['mes']} {datos['anio']}"
    inf = Informe()

    def n_ri(tipo):
        return int(ratio.loc[tipo, col_n]), _pct(ratio.loc[tipo, col_ri])

    inf.memo(datos["para"], datos["de"], datos["fecha"],
             f"Reporte de Control de Calidad – {periodo}",
             "Anexo A: Gráficos de CC" if anexo else "")

    # 1. Muestras evaluadas
    inf.titulo("Muestras de Control de Calidad Evaluados en este Reporte")
    inf.parrafo(f"Este reporte incluye la revisión de un total de {int(ratio.loc['Total Muestras', col_n])} "
                "muestras de blast holes, frentes de minado, duplicados, estándares, blancos finos y "
                "gruesos; analizados en el laboratorio SGS-Mina e introducidos en la base de datos desde "
                f"el {datos['periodo_desde']} hasta el día {datos['periodo_hasta']}. El resumen de muestras "
                "originales, controles insertados y check externo en este período se muestran en la "
                "Tabla 1:")
    inf.leyenda("Lotes Incluidos en el Reporte")
    ultima = len(datos["ratio"]) - 1
    inf.tabla(datos["ratio"], tamano=9, negrita_filas=(0, 1, ultima), gris_filas=(0, 1))

    # 2. Resultados
    inf.titulo("Resultados del Control de Calidad")
    inf.parrafo("Los gráficos se presentan en el Anexo A. Los duplicados fueron procesados por el "
                "método hiperbólico.")
    dup = rep["Duplicado"]
    textos_dup = {
        "DC": ("Duplicados de Campo", "duplicados de campo", 30,
               "En general, la varianza de muestreo está dentro de los límites aceptables",
               "Dup de Campo"),
        "RG": ("Duplicados Gruesos de Rechazo", "duplicados gruesos", 20,
               "La varianza de submuestreo durante preparación en laboratorio está dentro de los "
               "límites aceptables", "Rechazo Grueso"),
        "DP": ("Duplicados de Pulpa", "duplicados de pulpa", 10,
               "La varianza analítica está dentro de los límites aceptables", "Dup de Pulpa"),
    }
    tabla_dup = None
    for clave, (titulo, nombre, error_max, frase_ok, tipo_ratio) in textos_dup.items():
        inf.subtitulo(titulo)
        n, ri = n_ri(tipo_ratio)
        malos = elementos_con_clase(dup, f"Clase {clave}", "INACEPTABLE")
        if tabla_dup is None:
            tabla_dup = inf.n_tabla + 1
        inf.parrafo(
            f"En total, se revisaron los resultados de {n} {nombre}, lo que representa una tasa de "
            f"inserción de {ri}%. Se prepararon gráficos Max-Min para Cu, Ag, Mo, As, Zn y Fe "
            f"(Anexo A). {texto_tasas(dup, f'%Fallos {clave}')} (Tabla {tabla_dup}). "
            + (f"En este periodo {lista_es(malos)} {'presenta' if len(malos) == 1 else 'presentan'} "
               f"una tasa de error >{cfg.UMBRAL_FALLOS_DUPLICADO}% por ello se clasifica como "
               "inaceptable. " if malos else "En este periodo no se identificaron errores groseros. ")
            + f"Los resultados de los {nombre} se consideran "
            + (f"dentro del rango aceptable para todos los elementos estudiados excepto {lista_es(malos)} "
               if malos else "dentro del rango aceptable para todos los elementos estudiados ")
            + f"(más del 90 % de los pares deben tener un error relativo máximo de {error_max}%). "
            + (f"{frase_ok} a excepción de {lista_es(malos)}." if malos else f"{frase_ok}."),
            " ", ("PENDIENTE", "[Interpretación]"))
        if clave == "DC":
            inf.leyenda("Resumen de Duplicados")
            inf.tabla(dup)

    for clave, titulo, nombre, tipo_ratio in (("BG", "Blancos Gruesos", "blancos gruesos", "Blanco Grueso"),
                                              ("BF", "Blancos Finos", "blancos finos", "Blanco Fino")):
        inf.subtitulo(titulo)
        tabla_b = rep["Blanco Grueso" if clave == "BG" else "Blanco Fino"]
        n, ri = n_ri(tipo_ratio)
        malos = elementos_con_clase(tabla_b, "Clase", "INACEPTABLE")
        inf.parrafo(
            f"En total, se revisaron {n} {nombre}, lo que representa una tasa de inserción de {ri}%. "
            "Se prepararon gráficos de Blanco versus Muestra Precedente para todos los elementos "
            f"estudiados (Anexo A). {texto_tasas(tabla_b, '%Fallos')}, clasificándose como "
            + ("Aceptable" if not malos else f"Inaceptable para {lista_es(malos)}")
            + f". El resumen de {nombre} se muestra en la Tabla {inf.n_tabla + 1}. ",
            ("PENDIENTE", "[Interpretación del control de contaminación]"))
        inf.leyenda(f"Resumen de {titulo}")
        inf.tabla(tabla_b)

    # Estandares
    inf.subtitulo("Materiales de Referencia Certificados (MRC) o Estándares")
    n_std, ri_std = n_ri("STD")
    detalle = lista_es(f"{len(res.estandares[c])} muestras {conf['codigo']}"
                       for c, conf in res.conf_std.items())
    std = rep["Estandares por Analito"]
    referenciales = [f"{conf['codigo']} {cfg.ETIQUETA[e]} (CV {cv:.2f}%)"
                     for conf in res.conf_std.values()
                     for e, cv, _ in estandares.elementos_referenciales(conf)]
    cuestionables = elementos_con_clase(std, "Clase", "CUESTIONABLE", prefijo="Nombre\nSTD")
    inaceptables = elementos_con_clase(std, "Clase", "INACEPTABLE", prefijo="Nombre\nSTD")
    if not inaceptables and not cuestionables:
        exactitud = "En general, se observa una exactitud aceptable para todos los elementos."
    else:
        exactitud = "Se observa exactitud aceptable para los demás elementos"
        if cuestionables:
            exactitud += f", cuestionable para {lista_es(cuestionables)}"
        if inaceptables:
            exactitud += f" e inaceptable para {lista_es(inaceptables)}"
        exactitud += "."
    inf.parrafo(
        f"En total, se procesaron {n_std} muestras de MRCs" + (f" ({detalle})" if detalle else "")
        + f", lo que representa una tasa de inserción de {ri_std}%. Se prepararon gráficos de control "
        "para cada elemento y para cada tipo de MR certificado (Anexo A), según recomendación de Wood. "
        + (f"Los elementos cuyo valor esperado (BV) tiene CV mayor a {cfg.UMBRAL_CV_ESTANDAR}% "
           f"({lista_es(referenciales)}) se consideran valores referenciales: se grafican, pero "
           "no se evalúa su sesgo (ver Comentarios Adicionales). " if referenciales else "")
        + f"{exactitud} El resumen de la evaluación de los MRCs se presenta en la Tabla "
        f"{inf.n_tabla + 1}.")
    inf.leyenda("Resumen de los Materiales de Referencia")
    inf.tabla(std, col_color=1, colores=datos["colores"])
    inf.parrafo(f"Clasificación de la exactitud: aceptable: sesgo de 0 a ±{cfg.SESGO_ACEPTABLE}%; "
                f"cuestionable: sesgo de +{cfg.SESGO_ACEPTABLE} a +{cfg.SESGO_CUESTIONABLE}% o de "
                f"-{cfg.SESGO_ACEPTABLE}% a -{cfg.SESGO_CUESTIONABLE}%; inaceptable: sesgo superior "
                f"a {cfg.SESGO_CUESTIONABLE}% o inferior a -{cfg.SESGO_CUESTIONABLE}%. "
                + (f"REFERENCIAL: elemento con CV del valor esperado mayor a "
                   f"{cfg.UMBRAL_CV_ESTANDAR}%, presentado solo como referencia. "
                   if referenciales else ""),
                ("PENDIENTE", "[Interpretación]"))

    # Check externo
    inf.subtitulo("Check Externo")
    rma = datos.get("reporte_rma")
    if rma is not None and not rma.empty:
        n_ext = int(rma["RMA\n# Pares"].max())
        clase_rma = {c: elementos_con_clase(rma, "RMA\nClase", c)
                     for c in ("ACEPTABLE", "CUESTIONABLE", "INACEPTABLE")}
        con_outliers = elementos_con_clase(rma, "Outlier (ER>0.1)\nClase", "INACEPTABLE")
        fuera_5 = elementos_con_clase(rma, "Certimin\ndentro 5%\nSGS", "INACEPTABLE")
        general_mal = elementos_con_clase(rma, "Clase\nGeneral", "INACEPTABLE")
        partes = []
        for clase, texto in (("ACEPTABLE", "aceptable"), ("CUESTIONABLE", "cuestionable"),
                             ("INACEPTABLE", "inaceptable")):
            if clase_rma[clase]:
                partes.append(f"{texto} para {lista_es(clase_rma[clase])}")
        inf.parrafo(
            f"En total, se enviaron para control externo {n_ext} muestras al laboratorio secundario "
            "Certimin, como complemento a la evaluación de la exactitud del laboratorio SGS-Mina. Se "
            "realizó la evaluación de los datos resultantes mediante el método de RMA (Sinclair, 1999, "
            "citado por Long, 2005), parcialidades relativas (promedios) y errores relativos (ER) menor "
            "de 10%. Se prepararon las gráficas correspondientes para los elementos Cu, Ag, Mo, As y Zn "
            "(Anexo A). "
            + (f"Se identificaron más de {cfg.UMBRAL_OUTLIER_EXT}% de outliers con ER superiores al 10% "
               f"para {lista_es(con_outliers)}. " if con_outliers else
               "La proporción de outliers con ER superiores al 10% fue aceptable para todos los elementos. ")
            + f"Según el método RMA, el sesgo es {'; '.join(partes)}. "
            + ("Las parcialidades relativas fueron aceptables para todos los elementos"
               + (f", excepto {lista_es(fuera_5)}." if fuera_5 else ".")),
            " ", ("PENDIENTE", "[Interpretación de los sesgos y outliers]"))
        inf.parrafo(f"En conclusión, la exactitud del laboratorio SGS Toromocho es "
                    f"{texto_aceptabilidad(general_mal)}. El resumen de la evaluación de la exactitud "
                    f"mediante el Check Externo se presenta en la Tabla {inf.n_tabla + 1}.")
        inf.leyenda("Resumen de Control Externo")
        inf.tabla(rma)
    else:
        general_mal = None
        inf.pendiente("No se procesó el control externo (pestaña RMA) para este periodo. "
                      "Completar o eliminar esta sección.")

    # Acumulado
    inf.titulo("Análisis del acumulado del año")
    inf.parrafo("Se ha realizado una evaluación de la exactitud de SGS Toromocho, a partir de los "
                f"estándares analizados desde {datos['periodo_desde']} a {periodo}. Cuyos resultados "
                "son los siguientes:")
    inf.figuras(list(acumulado), columnas=1, ancho_cm=15, avance=avance)
    inf.pendiente("Comentarios sobre la evolución de los estándares en el año, cambios de "
                  "protocolo del laboratorio, etc.")

    # Comentarios adicionales
    inf.titulo("Comentarios Adicionales")
    inf.parrafo("Para los análisis de contaminación y precisión, de acuerdo con las recomendaciones de "
                "Wood (2021) se están utilizando los Límites Prácticos de Detección (LPD), calculados a "
                "partir de duplicados de pulpas de muestras que fueron enviadas al laboratorio SGS-Mina. "
                f"Los valores se muestran en la Tabla {inf.n_tabla + 1}.")
    inf.leyenda("Límites prácticos de detección")
    inf.tabla(tabla_limites_deteccion(), tamano=9)
    inf.parrafo("Para MRCs, se están utilizando las medias y desviaciones estándar históricas para el "
                "análisis de exactitud en los materiales de referencia, de acuerdo con la recomendación "
                "de Wood (2021).")
    inf.parrafo(f"{CRITERIO_CV} Los MRCs utilizados y la categoría de cada elemento se "
                f"presentan en la Tabla {inf.n_tabla + 1}.")
    inf.leyenda("Materiales de Referencia utilizados en el periodo")
    inf.tabla(tabla_estandares_usados(res.conf_std), col_color=0, colores=datos["colores"])

    # Conclusiones
    inf.titulo("Conclusiones y Recomendaciones")
    inf.parrafo(f"Como resultado del procesamiento de los datos de CC durante el mes de {periodo}, "
                "se concluye lo siguiente:")
    dc_mal = elementos_con_clase(dup, "Clase DC", "INACEPTABLE")
    sub_mal = sorted(set(elementos_con_clase(dup, "Clase RG", "INACEPTABLE")
                         + elementos_con_clase(dup, "Clase DP", "INACEPTABLE")),
                     key=lambda x: [cfg.ETIQUETA[e] for e in cfg.ELEMENTOS].index(x))
    bg_mal = elementos_con_clase(rep["Blanco Grueso"], "Clase", "INACEPTABLE")
    bf_mal = elementos_con_clase(rep["Blanco Fino"], "Clase", "INACEPTABLE")
    conclusiones = [
        "Los resultados de duplicados de campo indican que la precisión de muestreo es "
        f"{texto_aceptabilidad(dc_mal, 'todos los elementos evaluados en el periodo')}.",
        "Los resultados de duplicados gruesos y de pulpa indican que la precisión de submuestreo en "
        "preparación y análisis químicos es "
        f"{texto_aceptabilidad(sub_mal, 'todos los elementos')}.",
        "De resultados de blancos gruesos, se observa un control de contaminación durante "
        f"preparación {texto_aceptabilidad(bg_mal)}.",
        "De resultados de blancos finos, se tiene un control de contaminación durante análisis "
        f"químico {texto_aceptabilidad(bf_mal)}.",
        "A partir de resultados de sesgo de estándares, se concluye que la exactitud analítica es "
        + (texto_aceptabilidad(inaceptables + cuestionables, "todos los elementos estudiados")
           + (" (ver Tabla de Materiales de Referencia)." if inaceptables or cuestionables else "."))
        + (f" {lista_es(referenciales)} se evaluaron solo como valores referenciales por tener "
           f"CV del valor esperado mayor a {cfg.UMBRAL_CV_ESTANDAR}%." if referenciales else ""),
    ]
    if general_mal is not None:
        conclusiones.append("Del RMA, que evalúa la exactitud de SGS Toromocho a partir de muestras "
                            "check externos analizados por Certimin, se observa exactitud "
                            f"{texto_aceptabilidad(general_mal)}.")
    for texto in conclusiones:
        inf.parrafo(texto, " ", ("PENDIENTE", "[Comentario]"), vineta=True)
    inf.parrafo(("PENDIENTE", "[Agregar observaciones de QQ-Plot por rango de leyes y "
                              "recomendaciones]"), vineta=True)

    # Anexo A
    if anexo:
        inf.salto_pagina()
        inf.titulo("ANEXOS GRÁFICOS", centrado=True, tamano=12)
        inf.titulo("ANEXO A: Gráficos de Control de Calidad")
        for titulo, graficos in anexo:
            if graficos:
                inf.subtitulo(titulo)
                inf.figuras(graficos, avance=avance)
    inf.guardar(ruta)
    return ruta


def fechas_periodo(bd_total):
    """(desde, hasta) como dd.mm.aaaa usando la columna FECHA de la BD de leyes."""
    for columna in ("FECHA", "FECHA_REPORT", "FECHA DE ENVIO"):
        if columna in bd_total:
            fechas = pd.to_datetime(bd_total[columna], errors="coerce").dropna()
            if not fechas.empty:
                return fechas.min().strftime("%d.%m.%Y"), fechas.max().strftime("%d.%m.%Y")
    hoy = datetime.now().strftime("%d.%m.%Y")
    return hoy, hoy
