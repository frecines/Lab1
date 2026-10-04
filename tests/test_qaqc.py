# -*- coding: utf-8 -*-
"""Pruebas con datos sinteticos: python -m pytest tests"""
import os
import sys

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import qaqc_calc as calc  # noqa: E402
import qaqc_config as cfg  # noqa: E402
import qaqc_estandares as std  # noqa: E402

ESTANDARES = std.PREDETERMINADOS

# Cada bloque de 20 muestras: (posicion, CONTROLES)
PATRON = {0: "CAMPO ORIGIN", 3: "RECHA", 6: "PULPA", 8: "BLK FINO",
          11: "BLANCO", 13: "STD-13", 15: "STD-14", 17: "STD-15"}


def datos_sinteticos(n=400, inicio=1000, semilla=1):
    rng = np.random.default_rng(semilla)
    ids = np.arange(inicio, inicio + n)
    controles = [PATRON.get(i % 20) for i in range(n)]
    leyes = pd.DataFrame({
        "SAMPLE_ID": ids,
        "LAB_REPORT": [f"LR-{i // 50}" for i in range(n)],
        "CU": rng.lognormal(-1.5, 0.6, n), "AG": rng.lognormal(0.5, 0.6, n),
        "MO": rng.lognormal(-5, 0.6, n), "ARS": rng.lognormal(-4.5, 0.6, n),
        "ZN": rng.lognormal(-3.5, 0.6, n), "FE": rng.lognormal(0.5, 0.3, n),
        "CUOX": rng.lognormal(-3, 0.5, n),
    })
    for i, c in enumerate(controles):
        if c == "CAMPO ORIGIN":            # duplicado = muestra siguiente
            leyes.loc[i + 1, cfg.COLUMNA_BD.values()] = (
                leyes.loc[i, cfg.COLUMNA_BD.values()].to_numpy(float) * rng.normal(1, 0.08, 6))
        elif c in ("RECHA", "PULPA"):      # duplicado = fila de control
            leyes.loc[i, cfg.COLUMNA_BD.values()] = (
                leyes.loc[i - 1, cfg.COLUMNA_BD.values()].to_numpy(float) * rng.normal(1, 0.05, 6))
        elif c in ("BLK FINO", "BLANCO"):
            leyes.loc[i, cfg.COLUMNA_BD.values()] = [0.0005, 0.15, 0.0007, 0.001, 0.0008, 0.05]
        elif c and c.startswith("STD"):
            conf = ESTANDARES[c]
            for e in cfg.ELEMENTOS_STD:
                leyes.loc[i, cfg.COLUMNA_BD[e]] = rng.normal(conf["MH"][e], conf["DE"][e] * 1.5)
    leyes["CU"] = leyes["CU"].astype(object)
    leyes.loc[1, "CU"] = "<0.001"          # texto de limite de deteccion
    leyes.loc[2, "CU"] = "<0.0001"         # texto no listado -> x/2
    meses = ["ENERO", "FEBRERO", "MARZO", "ABRIL"]
    sec = pd.DataFrame({
        "SAMPLE_ID": ids, "CONTROLES": controles,
        "FECHA DE ENVIO": pd.Timestamp("2024-01-01"),
        "MES ": [meses[i * len(meses) // n] for i in range(n)],
        "Malla": [100 if i % 2 else 200 for i in range(n)],
        "Scheme": [f"R{i // 100}" for i in range(n)],
    })
    return leyes, sec


@pytest.fixture(scope="module")
def resultados():
    leyes, sec = datos_sinteticos()
    total, avisos = calc.unir_bases(leyes, sec, 1000, 1399)
    assert avisos == []
    return calc.ResultadosQAQC(total, ESTANDARES)


def test_a_numero():
    s = pd.Series(["<0.001", "0,5", ">500", "<0.0001", "abc", None, 2])
    r = calc.a_numero(s, cfg.REEMPLAZO_LD)
    assert r.iloc[0] == 0.0005 and r.iloc[1] == 0.5 and r.iloc[2] == 501
    assert r.iloc[3] == pytest.approx(0.00005)
    assert np.isnan(r.iloc[4]) and np.isnan(r.iloc[5]) and r.iloc[6] == 2


def test_filtros():
    leyes, sec = datos_sinteticos()
    total, _ = calc.unir_bases(leyes, sec, 1000, 1399, malla="100", reporte="R1")
    assert len(total) == 50
    assert set(total["Malla"]) == {100} and set(total["Scheme"]) == {"R1"}
    with pytest.raises(ValueError):
        calc.unir_bases(leyes, sec, 5000, 6000)
    with pytest.raises(ValueError, match="LAB_REPORT"):
        calc.unir_bases(leyes.drop(columns="LAB_REPORT"), sec, 1000, 1399)


def test_pareo_duplicados(resultados):
    dc, rg = resultados.duplicados["DC"], resultados.duplicados["RG"]
    assert len(dc) == 20 and len(rg) == 20
    assert (dc["SAMPLE_DUP"] == dc["SAMPLE_ORIG"] + 1).all()
    assert (rg["SAMPLE_ORIG"] == rg["SAMPLE_ID"] - 1).all()
    total = resultados.bd_total.set_index("SAMPLE_ID")
    fila = rg.iloc[0]
    assert fila["CU_O"] == total.loc[fila["SAMPLE_ORIG"], "CU"]
    assert fila["CU_D"] == total.loc[fila["SAMPLE_ID"], "CU"]
    # duplicados con ~5-8% de dispersion deben pasar casi siempre
    assert calc.tasa(rg, "CU_ERR") < 20


def test_blancos(resultados):
    bf, bg = resultados.blancos["BF"], resultados.blancos["BG"]
    assert (bf["SAMPLE_PREC"] == bf["SAMPLE_ID"] - 1).all()
    assert (bg["SAMPLE_PREC"] == bg["SAMPLE_ID"] - 2).all()
    assert bf["CU_MAV"].iloc[0] == pytest.approx(3 * 0.001)
    assert bg["FE_MAV"].iloc[0] == pytest.approx(5 * 0.127)
    assert calc.tasa(bf, "CU_ERR") == 0


def test_reportes(resultados):
    r = resultados.reportes
    assert list(r["Duplicado"].columns[:4]) == ["Analito", "N° Muestras DC",
                                                "%Fallos DC", "Clase DC"]
    std = r["Estandares"]
    assert len(std) == 15
    no_cert = std[std["Media\nCertificada\n(BV)"].isna()]
    assert set(no_cert["Clase"]) == {"No Certificado"}
    assert len(no_cert) == 3   # Ag y As en CHNB13, Ag en CHNM14
    assert list(r["Estandares por Analito"]["Analito"][:3]) == ["Cu %"] * 3
    # La lista de errores de Duplicado de Campo debe venir de DC (antes venia de Pulpa)
    dc_ids = set(resultados.duplicados["DC"]["SAMPLE_ORIG"])
    assert set(r["ErrorDuplicadoCampo"]["SAMPLE_ORIG"]) <= dc_ids
    assert "Ubicacion Error" in r["TotalError"].columns


def test_clasificar_sesgo_limites():
    assert calc.clasificar_sesgo(5.0) == "ACEPTABLE"       # antes: INACEPTABLE
    assert calc.clasificar_sesgo(-5.0) == "ACEPTABLE"
    assert calc.clasificar_sesgo(7) == "CUESTIONABLE"
    assert calc.clasificar_sesgo(10.5) == "INACEPTABLE"
    assert calc.clasificar_sesgo(1, certificado=False) == "No Certificado"


def test_exportar(tmp_path, resultados):
    ruta = tmp_path / "Reporte_QAQC.xlsx"
    calc.exportar_excel(resultados.reportes, ruta)
    hojas = pd.ExcelFile(ruta).sheet_names
    assert "Estandares" in hojas and "TotalError" in hojas


def _excel_externo(tmp_path, ids, rng):
    """Archivo con el formato de Certimin: 9 filas de cabecera, fila de
    codigos, fila de nombres, 3 filas extra y luego los datos."""
    nombres = ["SAMPLE", "Cu", "Ag", "Mo", "As", "Zn", "CuSolH+", "CuScitrico"]
    codigos = ["X", "G0172MCR1", "G0125MCR1", "A", "B", "C", "D", "E"]
    filas = [["basura"] * 8 for _ in range(9)] + [codigos, nombres] + [[None] * 8] * 3
    for i in ids:
        filas.append([str(i), *np.round(rng.lognormal(-1.5, 0.5, 7), 4)])
    filas[-1][1] = "<0.001"
    ruta = tmp_path / "externo.xlsx"
    pd.DataFrame(filas).to_excel(ruta, header=False, index=False)
    return ruta


def test_externo(tmp_path, resultados):
    rng = np.random.default_rng(3)
    ruta = _excel_externo(tmp_path, range(1005, 1100, 3), rng)
    cert = calc.leer_reporte_externo(ruta)
    rma = calc.procesar_externo(cert, resultados.bd_total, 1000, 1399)
    assert len(rma) == 32 and rma["Cu_ct"].iloc[-1] == 0.0005
    rep = calc.reporte_externo(rma)
    assert list(rep["Analito"]) == ["Cu %", "Ag gr/tn", "Mo %", "As %", "Zn %", "CuSS %"]
    p = calc.pares_externo(rma, "CU")
    m, b = calc.recta_rma(p["X"], p["Y"])
    assert np.isfinite(m) and np.isfinite(b)


def test_ratio(resultados):
    df = calc.ratio_insercion(resultados.bd_total, 7, "2024", "Enero")
    assert df.iloc[0, 1] == 400
    assert df.set_index("Año/Mes").loc["STD", "2024"] == 60
    assert df.iloc[-1, 1] == 20 * len(PATRON) + 7


def test_graficos(resultados):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    import qaqc_plots as plots

    def dibujar(funcion, **kw):
        fig = Figure()
        funcion(fig, **kw)
        fig.canvas.draw()

    for clave in cfg.BLANCOS:
        df = resultados.blancos[clave]
        dibujar(plots.grafico_blanco, df=df, elemento="CU", titulo="t",
                tasa_error=calc.tasa(df, "CU_ERR"), meses=True)
        dibujar(plots.grafico_precedente, df=df, elemento="FE", titulo="t")
    for clave, conf in cfg.DUPLICADOS.items():
        df = resultados.duplicados[clave]
        dibujar(plots.grafico_duplicado, df=df, elemento="AG", conf=conf, tasa_error=1.0)
        dibujar(plots.grafico_qq, x=df["AG_O"], y=df["AG_D"], titulo="t", xlabel="x",
                ylabel="y", nombre_x="a", nombre_y="b")
        dibujar(plots.grafico_rd, df=df, col_prom="AG_PROM", col_rd="AG_RD", titulo="t",
                etiqueta="Ag", id_col="SAMPLE_ORIG")
    for clave, conf in resultados.conf_std.items():
        dibujar(plots.grafico_estandar, df=resultados.estandares[clave], elemento="AS",
                conf=conf, meses=True)
    vacio = resultados.blancos["BF"].iloc[0:0]
    dibujar(plots.grafico_blanco, df=vacio, elemento="CU", titulo="t", tasa_error=np.nan)


# ------------------------------------------------- parametros de estandares
EXCEL_ESTANDARES = ("/root/.claude/uploads/06f4bf4c-f7a3-591b-9e7b-da240b7d8c1c/"
                    "e98e23b9-Estandares_2026.xlsx")


def test_solo_estandares_con_muestras(resultados):
    # Los datos sinteticos tienen STD-13/14/15; STD-16/17/18 no se reportan
    assert set(resultados.conf_std) == {"STD-13", "STD-14", "STD-15"}
    assert "ErrorCHNB13" in resultados.reportes
    tabla = calc.errores_estandares(resultados)
    assert list(tabla.columns) == ["STD", "SAMPLE_ID", "LAB_REPORT", "OBS"]


@pytest.mark.skipif(not os.path.exists(EXCEL_ESTANDARES), reason="Excel de ejemplo no disponible")
def test_leer_excel_estandares():
    leidos = std.leer_excel(EXCEL_ESTANDARES, std.PREDETERMINADOS)
    assert list(leidos) == ["STD-16", "STD-17", "STD-18"]
    for control, conf in leidos.items():
        for campo in std.CAMPOS:
            for e in cfg.ELEMENTOS_STD:
                assert conf[campo][e] == pytest.approx(ESTANDARES[control][campo][e], rel=1e-12)
        assert conf["CV"]["AG"] > 5   # Ag resaltado en amarillo en el Excel
        assert conf["codigo"] == ESTANDARES[control]["codigo"]  # se conserva


def test_exportar_e_importar(tmp_path):
    ruta = tmp_path / "est.xlsx"
    std.exportar_excel(ESTANDARES, ruta)
    leidos = std.leer_excel(ruta)
    assert list(leidos) == list(ESTANDARES)
    for control, conf in ESTANDARES.items():
        assert leidos[control]["codigo"] == conf["codigo"]
        assert leidos[control]["titulo"] == conf["titulo"]
        for campo in std.CAMPOS:   # Excel guarda 15 cifras significativas
            assert leidos[control].get(campo, {}) == pytest.approx(conf.get(campo, {}), rel=1e-13)


def test_parametros_guardar_e_historial(tmp_path):
    p = std.Parametros(str(tmp_path))
    assert p.clave_correcta(std.CLAVE_INICIAL)
    nuevos = std._copia(p.estandares)
    nuevos["STD-16"]["MH"]["CU"] = 0.25
    nuevos["STD-19"] = dict(codigo="X19", titulo="Prueba", BV={},
                            MH={e: 1.0 for e in cfg.ELEMENTOS_STD},
                            DE={e: 0.1 for e in cfg.ELEMENTOS_STD})
    cambios = p.guardar(nuevos, std.hash_clave("nueva"))
    assert ("STD-16", "Cu", "MH", repr(0.2454732824427479), repr(0.25)) in cambios
    assert any(c[0] == "STD-19" and c[4] == "agregado" for c in cambios)
    p2 = std.Parametros(str(tmp_path))          # se vuelve a leer del archivo
    assert p2.estandares["STD-16"]["MH"]["CU"] == 0.25 and "STD-19" in p2.estandares
    assert p2.clave_correcta("nueva") and not p2.clave_correcta(std.CLAVE_INICIAL)
    historial = (tmp_path / std.HISTORIAL).read_text(encoding="utf-8-sig")
    assert "STD-16;Cu;MH" in historial and "Contrasena" in historial


def test_validacion_estandares():
    malo = std._copia(ESTANDARES)
    del malo["STD-18"]["DE"]["ZN"]
    with pytest.raises(ValueError, match="STD-18 Zn: falta DE"):
        std.validar(malo)


def test_archivo_corrupto_usa_fabrica(tmp_path):
    (tmp_path / std.ARCHIVO).write_text("{no es json", encoding="utf-8")
    p = std.Parametros(str(tmp_path))
    assert p.aviso and "STD-18" in p.estandares


# ----------------------------------------------------------- CV > 5% e informe
def test_cv_omite_graficos():
    graficar, omitidos = std.elementos_graficables(ESTANDARES["STD-16"])
    assert graficar == ["CU", "MO", "AS", "ZN"]
    assert [e for e, _ in omitidos] == ["AG"]
    # Estandares antiguos sin CV cargado: se grafican todos
    assert std.elementos_graficables(ESTANDARES["STD-13"])[0] == cfg.ELEMENTOS_STD
    for control in ("STD-17", "STD-18"):
        assert "AG" not in std.elementos_graficables(ESTANDARES[control])[0]


def test_textos_informe():
    import qaqc_informe as inf
    assert inf.lista_es(["Cu", "Ag", "Zn"]) == "Cu, Ag y Zn"
    rep = pd.DataFrame({"Analito": ["Cu %", "Mo %", "Zn %"], "%F": [2.721, 0.0, 10.884]})
    assert inf.texto_tasas(rep, "%F") == "Las tasas de error oscilaron entre 0% para Mo y 10.88% para Zn"
    rep["%F"] = 0.0
    assert inf.texto_tasas(rep, "%F") == "La tasa de error fue de 0% para todos los elementos"


def test_generar_informe(tmp_path, resultados):
    import functools
    import docx
    import qaqc_informe as inf
    import qaqc_plots as plots
    ratio = calc.ratio_insercion(resultados.bd_total, 0, "2025", "ENERO")
    datos = dict(para="Juan Ramos", de="Equipo Geología Mina", fecha="22/02/2025",
                 mes="Enero", anio="2025", periodo_desde="26.12.2024", periodo_hasta="25.01.2025",
                 res=resultados, ratio=ratio, reporte_rma=None,
                 colores=std.colores(ESTANDARES))
    df = resultados.duplicados["DC"]
    anexo = [("Duplicado de Campo", [("Cu", "x", (8, 8), functools.partial(
        plots.grafico_duplicado, df=df, elemento="CU", conf=cfg.DUPLICADOS["DC"], tasa_error=1.0))])]
    ruta = inf.generar(str(tmp_path / "informe.docx"), datos, anexo)
    documento = docx.Document(ruta)
    texto = "\n".join(p.text for p in documento.paragraphs)
    assert "Reporte de Control de Calidad – Enero 2025" in "\n".join(
        c.text for c in documento.tables[0].columns[1].cells)
    assert "un total de 400 muestras" in texto
    assert "20 duplicados de campo" in texto and "ANEXO A" in texto
    assert len(documento.inline_shapes) == 1
    assert len(documento.tables) >= 9   # memo + 7 tablas + grilla del anexo
